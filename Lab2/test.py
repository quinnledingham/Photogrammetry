from docx import Document

def array_to_word_table(array, filename="output.docx"):
    doc = Document()
    table = doc.add_table(rows=len(array), cols=len(array[0]))

    for i, row in enumerate(array):
        for j, cell in enumerate(row):
            table.cell(i, j).text = str(cell)

    doc.save(filename)
    print(f"Word document '{filename}' created successfully.")

# Example usage:
data = [
    ["Name", "Age", "City"],
    ["Alice", 30, "New York"],
    ["Bob", 25, "Los Angeles"],
    ["Charlie", 35, "Chicago"]
]

array_to_word_table(data)
